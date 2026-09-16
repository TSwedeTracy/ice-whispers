import json
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\trasa\Elgiganten Cloud\TSwede Biz\TSwedeBiz Test Session Tllda.odt")
OUTPUT = Path(r"C:\Users\trasa\Elgiganten Cloud\Affärsideer och projekt\Kortlekar\Magic Runes\ICE Whispers APP\TSwedeBiz - Samma text ren layout.docx")

NAVY = "23364D"
TEAL = "2B7A78"
PLUM = "71557A"
INK = "263238"
MUTED = "66737B"
PALE_BLUE = "EAF1F7"
PALE_TEAL = "E7F3F1"
PALE_PLUM = "F1ECF4"
PALE_GOLD = "FFF5DA"


def paragraphs_from_odt(path):
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("content.xml"))
    return ["".join(e.itertext()) for e in root.iter() if e.tag.endswith("}p") or e.tag.endswith("}h")]


def extract_json(lines, start_marker, end_marker=None):
    start = lines.index(start_marker) + 1
    end = lines.index(end_marker) if end_marker else len(lines)
    return json.loads("\n".join(lines[start:end]).strip())


def set_font(run, size=10.5, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def fill_paragraph(p, fill):
    ppr = p._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    p.paragraph_format.page_break_before = page_break
    p.add_run(text)
    return p


def add_value(doc, value, fill=None, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(6)
    if bullet:
        p.paragraph_format.left_indent = Inches(0.48)
        p.paragraph_format.first_line_indent = Inches(-0.23)
    else:
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.right_indent = Inches(0.10)
    if fill:
        fill_paragraph(p, fill)
    r = p.add_run(str(value))
    set_font(r)


def render_object(doc, obj, level=1, alternating=0):
    if isinstance(obj, dict):
        for key, value in obj.items():
            if isinstance(value, (dict, list)):
                add_heading(doc, key, min(level, 3))
                render_object(doc, value, min(level + 1, 3), alternating + 1)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.left_indent = Inches(0.12)
                fill_paragraph(p, PALE_TEAL if alternating % 2 == 0 else PALE_BLUE)
                r = p.add_run(key)
                set_font(r, size=9.3, color=TEAL, bold=True)
                r = p.add_run("\n" + str(value))
                set_font(r, size=10.4)
    elif isinstance(obj, list):
        for index, item in enumerate(obj):
            if isinstance(item, dict):
                if index:
                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.space_after = Pt(2)
                render_object(doc, item, level, alternating + index)
            else:
                add_value(doc, item, bullet=True)
    else:
        add_value(doc, obj)


lines = paragraphs_from_odt(SOURCE)
free_report = extract_json(lines, "=== FREE REPORT ===", "=== DEEP DIVE CHAT TRANSCRIPT ===")
deep_report = extract_json(lines, "=== DEEP DIVE REPORT ===", "=== MY BUSINESS BLUEPRINT ===")
blueprint = extract_json(lines, "=== MY BUSINESS BLUEPRINT ===")

chat_start = lines.index("=== DEEP DIVE CHAT TRANSCRIPT ===") + 1
chat_end = lines.index("=== DEEP DIVE REPORT ===")
chat = lines[chat_start:chat_end]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.8)
sec.left_margin = Inches(0.95)
sec.right_margin = Inches(0.95)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.13

for name, size, color, before, after in (
    ("Heading 1", 17, NAVY, 16, 8),
    ("Heading 2", 13.5, TEAL, 12, 6),
    ("Heading 3", 11.5, PLUM, 9, 4),
):
    style = doc.styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Original opening text only.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(12)
r = p.add_run(lines[0])
set_font(r, size=23, color=NAVY, bold=True)
for text in lines[1:5]:
    if text:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_font(r, size=11, color=INK, bold=text == "GOALS:")

add_heading(doc, "FREE REPORT", 1, page_break=True)
render_object(doc, free_report)

add_heading(doc, "DEEP DIVE CHAT TRANSCRIPT", 1, page_break=True)
for text in chat:
    p = doc.add_paragraph()
    if not text:
        p.paragraph_format.space_after = Pt(3)
        continue
    is_owner = text.startswith("Owner:")
    is_ai = text.startswith("AI:")
    p.paragraph_format.left_indent = Inches(0.48 if is_owner else 0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    if is_owner:
        fill_paragraph(p, PALE_PLUM)
    elif is_ai:
        fill_paragraph(p, PALE_TEAL)
    r = p.add_run(text)
    set_font(r, color=PLUM if is_owner else NAVY if is_ai else MUTED, bold=is_owner)

add_heading(doc, "DEEP DIVE REPORT", 1, page_break=True)
render_object(doc, deep_report)

add_heading(doc, "MY BUSINESS BLUEPRINT", 1, page_break=True)
render_object(doc, blueprint)

doc.core_properties.title = lines[0]
doc.core_properties.subject = "Same wording with clean reading layout"
doc.save(OUTPUT)
print(OUTPUT)
