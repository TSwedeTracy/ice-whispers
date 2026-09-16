import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\trasa\Elgiganten Cloud\TSwede Biz\TSwedeBiz Test Session Tllda.odt")
OUTPUT = Path(r"C:\Users\trasa\Elgiganten Cloud\Affärsideer och projekt\Kortlekar\Magic Runes\ICE Whispers APP\TSwedeBiz - Ordagrann formaterad version.docx")

NAVY = "24364B"
TEAL = "2A7F7F"
PURPLE = "6F5A8A"
INK = "263238"
MUTED = "64717A"
PALE_BLUE = "EAF1F8"
PALE_PURPLE = "F1EDF7"
PALE_TEAL = "EAF5F3"


def get_paragraphs(path):
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("content.xml"))
    return ["".join(e.itertext()) for e in root.iter() if e.tag.endswith("}p") or e.tag.endswith("}h")]


def set_font(run, name="Aptos", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def paragraph_fill(paragraph, color):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), color)


lines = get_paragraphs(SOURCE)
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

section_markers = {
    "=== FREE REPORT ===",
    "=== DEEP DIVE CHAT TRANSCRIPT ===",
    "=== DEEP DIVE REPORT ===",
    "=== MY BUSINESS BLUEPRINT ===",
}

for index, text in enumerate(lines):
    p = doc.add_paragraph()
    if index == 0:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text)
        set_font(r, size=23, color=NAVY, bold=True)
    elif text in section_markers:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(12)
        paragraph_fill(p, PALE_BLUE)
        r = p.add_run(text)
        set_font(r, size=16, color=NAVY, bold=True)
    elif text == "GOALS:":
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        set_font(r, size=13, color=TEAL, bold=True)
    elif text.startswith("AI: "):
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(3)
        paragraph_fill(p, PALE_TEAL)
        r = p.add_run(text)
        set_font(r, size=10.5, color=NAVY)
    elif text.startswith("Owner: "):
        p.paragraph_format.left_indent = Inches(0.55)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        paragraph_fill(p, PALE_PURPLE)
        r = p.add_run(text)
        set_font(r, size=10.5, color=PURPLE, bold=True)
    elif text in ("{", "}", "[", "]", "},", "],") or text.startswith('"'):
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, name="Aptos Mono", size=8.7, color=INK)
    elif not text:
        p.paragraph_format.space_after = Pt(4)
    else:
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_font(r)

doc.core_properties.title = lines[0] if lines else "TSwedeBiz Test Session"
doc.core_properties.subject = "Ordagrann formaterad version av originalet"
doc.save(OUTPUT)
print(OUTPUT)
